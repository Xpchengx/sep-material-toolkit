# -*- coding: utf-8 -*-
"""从 4 个本地源文件抽取统一明细表，输出 CSV + 汇总报告。"""
import os, io, csv, json, re, datetime
from openpyxl import load_workbook

# 台账根目录：优先取环境变量 LEDGER_DIR，默认用仓库内 data/ 目录
# （本地使用时把 LEDGER_DIR 指向你的台账所在目录）
_REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.environ.get('LEDGER_DIR') or os.path.join(_REPO, 'data')

# 源文件路径：**优先读 pick_sources.py 生成的 sources.json**（它会自动挑出
# 每个集成商最新的一份），没有该文件时回落到下面的硬编码路径。
# 这样新台账到了不用改代码，跑一次 pick_sources.py 就行。
_SRC_CFG = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'sources.json')
ROLES = {}          # 形如 {'A': '集成商A', 'B': '集成商B', 'C': '集成商C'}
if os.path.exists(_SRC_CFG):
    with io.open(_SRC_CFG, encoding='utf-8') as _f:
        _cfg = json.load(_f)
    SRC = _cfg.get('sources') or {}
    ROLES = _cfg.get('roles') or {}
    SRC_FROM = f'sources.json（{len(SRC)} 项，{_cfg.get("generatedAt", "?")}）'
else:
    SRC = {
        "集成商C": BASE + r"\集成商\集成商C\集成商C库存(9月8日).xlsx",
        "集成商A": BASE + r"\集成商\集成商A\上海集成商A材料库存(9月1日）.xlsx",
        "集成商B": BASE + r"\集成商\集成商B\集成商B-材料库存2026-8-26.xlsx",
        "采购": BASE + r"\采购与设计量与到货量.xlsx",
    }
    SRC_FROM = '脚本内硬编码路径（建议改用 pick_sources.py）'

# ---------------------------------------------------------------------------
# 三家台账的**格式**各不相同，下面按「角色」分支，而不是按名字。
# 角色到实际名称的对应由 sources.json 的 roles 给出；没给就沿用仓库里的脱敏名。
#   角色 A：无「月份」列；盘点表的 账面/实盘/差异 在 v[3..5]
#   角色 B：有「月份」列；盘点表列位在 v[6..8]
#   角色 C：有「设计明细」表；盘点表列位在 v[4..6]
# 这样换集成商、或把标签改成脱敏名，都不用改代码。
# ---------------------------------------------------------------------------
ROLE_A = ROLES.get('A') or "集成商A"
ROLE_B = ROLES.get('B') or "集成商B"
ROLE_C = ROLES.get('C') or "集成商C"
OUT = r"C:\Users\qazws\WorkBuddy\2026-09-09-09-19-37\material-system\out"
os.makedirs(OUT, exist_ok=True)

report = []
def log(*a):
    s = " ".join(str(x) for x in a)
    report.append(s)

log(f"源文件来源: {SRC_FROM}")


def txt(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return re.sub(r"\s+", " ", v).strip()
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", "")
    m = re.match(r"^-?\d+(\.\d+)?", s)
    return float(m.group()) if m else None


EPOCH = datetime.date(1899, 12, 30)


def dstr(v):
    """日期：datetime 直接用；Excel 序列号(20000~60000)换算。"""
    if v is None or v == "":
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    n = num(v)
    if n is not None and 20000 <= n <= 60000:
        try:
            return (EPOCH + datetime.timedelta(days=int(n))).strftime("%Y-%m-%d")
        except Exception:
            pass
    return txt(v)


def rows_of(ws, maxr=None):
    """迭代非全空行，返回 (rownum, [values])；用连续空行截断。"""
    blank = 0
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=maxr, values_only=True), start=1):
        vals = [txt(c) for c in row]
        while vals and vals[-1] == "":
            vals.pop()
        if not any(vals):
            blank += 1
            if blank >= 30:
                break
            continue
        blank = 0
        yield i, vals


def w(name, header, data):
    p = os.path.join(OUT, name + ".csv")
    with io.open(p, "w", encoding="utf-8-sig", newline="") as f:
        cw = csv.writer(f)
        cw.writerow(header)
        cw.writerows(data)
    log(f"  -> {name}.csv  {len(data)} 行")
    return len(data)


# ==================== 1. 三家集成商 ====================
inbound, outbound, returns, stock, invent = [], [], [], [], []
stock_seen = set()

for tag, path in [(k, SRC[k]) for k in [ROLE_C, ROLE_A, ROLE_B]]:
    if not path or not os.path.exists(path):
        log(f"\n## {tag}  !! 源文件不存在，跳过：{path}")
        continue
    log(f"\n## {tag}")
    wb = load_workbook(path, read_only=True, data_only=True)

    # --- 汇总 ---
    if "汇总" in wb.sheetnames:
        ws = wb["汇总"]
        n = 0
        for rn, vals in rows_of(ws):
            if rn <= 2:
                continue
            name = vals[1] if len(vals) > 1 else ""
            model = vals[2] if len(vals) > 2 else ""
            unit = vals[3] if len(vals) > 3 else ""
            rest = [v for v in vals[4:]]
            if not model and not name:
                continue
            # 末列是库存；倒数第2是出库；倒数第3是入库
            kucun = num(rest[-1]) if rest else None
            chuku = num(rest[-2]) if len(rest) > 1 else None
            ruku = num(rest[-3]) if len(rest) > 2 else None
            # 集成商B多一列「2025年结存」= 期初库存
            qichu = num(rest[-4]) if len(rest) > 3 else None
            stock.append([tag, name, model, unit, qichu, ruku, chuku, kucun])
            stock_seen.add((tag, model, name, unit))
            n += 1
        log(f"  汇总: {n} 行")

    # --- 入库明细 ---
    if "入库明细" in wb.sheetnames:
        ws = wb["入库明细"]
        n = 0
        for rn, vals in rows_of(ws):
            if rn == 1:
                continue
            vals = vals + [""] * 8
            # 角色 B 多一列「月份」：日期|月份|型号|数量|单位|材料来源|项目|备注
            if tag == ROLE_B:
                d, mon, model, qty, unit, frm, proj, rmk = vals[0], vals[1], vals[2], num(vals[3]), vals[4], vals[5], vals[6], vals[7]
            else:
                d, mon, model, qty, unit, frm, proj, rmk = vals[0], "", vals[1], num(vals[2]), vals[3], vals[4], vals[5], vals[6]
            if not model and not qty:
                continue
            inbound.append([tag, dstr(d), mon, model, qty, unit, frm, proj, rmk])
            n += 1
        log(f"  入库明细: {n} 行")

    # --- 出库明细 ---
    if "出库明细" in wb.sheetnames:
        ws = wb["出库明细"]
        n = 0
        for rn, vals in rows_of(ws):
            if rn == 1:
                continue
            vals = vals + [""] * 11
            if tag == ROLE_B:
                mon, d, model, qty, proj, to, who, rmk, yf, tb = (
                    vals[0], vals[1], vals[2], num(vals[3]), vals[4], vals[5], vals[6], vals[7], vals[8], vals[9])
            else:
                d, model, qty, proj, to, who, rmk, yf, rmk2, tb = (
                    vals[0], vals[1], num(vals[2]), vals[3], vals[4], vals[5], vals[6], vals[7], vals[8], vals[9])
                mon = ""
            if not model and not qty:
                continue
            outbound.append([tag, mon, dstr(d), model, qty, proj, to, who, rmk, yf, tb])
            n += 1
        log(f"  出库明细: {n} 行")

    # --- 退料/废料 ---
    for sn, kind in [("退料单", "退料单"), ("站点废料明细", "废料")]:
        if sn in wb.sheetnames:
            ws = wb[sn]
            n = 0
            for rn, vals in rows_of(ws):
                if rn <= 2:
                    continue
                vals = vals + [""] * 8
                # 站点废料明细: 日期|型号|数量|单位|项目|站点|退料人员|备注
                returns.append([tag, kind, dstr(vals[0]), vals[1], num(vals[2]), vals[3], vals[4], vals[5], vals[6], vals[7]])
                n += 1
            log(f"  {sn}: {n} 行")

    # --- 盘点表（取表头为 序号|物料名称|单位|... 的月表）---
    for sn in wb.sheetnames:
        if "盘" not in sn:
            continue
        ws = wb[sn]
        hdr = None
        for rn, vals in rows_of(ws):
            j = " ".join(vals)
            if "物料名称" in j and ("数量" in j or "结存" in j):
                hdr = rn
                break
            if rn > 6:
                break
        if hdr is None:
            continue
        n = 0
        for rn, vals in rows_of(ws):
            if rn <= hdr:
                continue
            v = vals + [""] * 10
            nm, model, unit, book, real, diff = v[1], "", v[2], None, None, None
            if tag == ROLE_C:
                model, unit = v[2], v[3]
                book, real, diff = num(v[4]), num(v[5]), num(v[6])
            elif tag == ROLE_A:
                book, real, diff = num(v[3]), num(v[4]), num(v[5])
            else:
                book, real, diff = num(v[6]), num(v[7]), num(v[8])
            if not nm:
                continue
            invent.append([tag, sn, nm, model, unit, book, real, diff, ""])
            n += 1
        log(f"  {sn}: {n} 行")
    wb.close()

log("")
n_in = w("入库明细", ["集成商", "日期", "月份", "型号", "数量", "单位", "材料来源", "项目", "备注"], inbound)
n_out = w("出库明细", ["集成商", "月份", "日期", "型号", "数量", "项目", "材料去向", "领料人", "备注", "应发未发", "退废料"], outbound)
n_ret = w("退料明细", ["集成商", "类型", "日期", "型号", "数量", "单位", "项目", "站点", "退料人员", "备注"], returns)
n_stk = w("集成商库存汇总", ["集成商", "名称", "型号", "单位", "期初库存", "入库", "出库", "库存"], stock)
n_inv = w("盘点记录", ["集成商", "盘点表", "物料名称", "型号", "单位", "账面数量", "实盘数量", "差异", "备注"], invent)

# ==================== 2. 角色 C 的设计明细 ====================
log(f"\n## {ROLE_C} 设计明细")
# 站点 -> 项目 映射：取自出库明细（每行都有站点与项目，比设计明细的项目列可靠）
_tmp = {}
for r in outbound:
    proj, to = r[5], r[6]
    if to and proj and to.strip() not in ("-", "无", ""):
        _tmp.setdefault(to.strip(), {})[proj.strip()] = _tmp.setdefault(to.strip(), {}).get(proj.strip(), 0) + 1
site2proj = {k: max(v.items(), key=lambda x: x[1])[0] for k, v in _tmp.items()}
log(f"  站点→项目 映射（来自出库明细）: {len(site2proj)} 条")

wb = load_workbook(SRC[ROLE_C], read_only=True, data_only=True)
ws = wb["设计明细"]
design = []
cur_site, cur_proj = "", ""
for rn, vals in rows_of(ws):
    if rn == 1:
        continue
    vals = vals + [""] * 6
    site, model, qty, proj, rmk = vals[1], vals[2], num(vals[3]), vals[4], vals[5]
    if site:
        # 新站点块：只取本行自带的项目，不继承上一块（避免串项目）
        cur_site = site.strip()
        cur_proj = proj.strip() if proj else ""
    elif proj and not cur_proj:
        cur_proj = proj.strip()
    if not model:
        continue
    # 项目优先取站点映射，其次用设计明细自带的
    proj_final = site2proj.get(cur_site) or cur_proj
    design.append([ROLE_C, proj_final, cur_site, model, qty, rmk])
wb.close()
n_des = w("站点设计量明细", ["集成商", "项目", "站点", "型号", "设计数量", "备注"], design)

# ==================== 3. 采购表 ====================
log("\n## 采购表 逐项目")
wb = load_workbook(SRC["采购"], read_only=True, data_only=True)
purchase = []
proj_meta = []
for sn in wb.sheetnames:
    ws = wb[sn]
    allrows = list(rows_of(ws))
    if len(allrows) < 5:
        log(f"  [{sn}] 行数不足，跳过")
        continue
    # 表头：r1 大组，r2 供应商名，r3 订单号，r4 订单数量/接收数量
    r2 = allrows[1][1] if len(allrows) > 1 else []
    r3 = allrows[2][1] if len(allrows) > 2 else []
    r4 = allrows[3][1] if len(allrows) > 3 else []
    # 找 r4 中的「订单数量」列（1-based = 索引+1）
    pairs = []
    c = 1
    while c <= len(r4):
        if "订单数量" in r4[c - 1]:
            sup = ""
            # 供应商名在 r2，可能合并 2 列
            for cc in range(c, min(c + 2, len(r2) + 1)):
                if r2[cc - 1]:
                    sup = r2[cc - 1]
                    break
            if not sup:
                for cc in range(c, 0, -1):
                    if cc - 1 < len(r2) and r2[cc - 1]:
                        sup = r2[cc - 1]
                        break
            oid = r3[c - 1] if c - 1 < len(r3) else ""
            pairs.append((c, sup, oid))
        c += 1
    cb = 5  # 设计量列
    n = 0
    n_dz = 0
    for rn, vals in allrows:
        if rn < 5:
            continue
        v = vals + [""] * (max([10] + [p[0] + 1 for p in pairs]))
        code, name, unit = v[1], v[2], v[3]
        dz, cg1, cg2, cgh, diff = num(v[4]), num(v[5]), num(v[6]), num(v[7]), num(v[8])
        if not name and not code:
            continue
        # 采购汇总/差额是公式，无缓存值时自行推算
        if cgh is None and (cg1 is not None or cg2 is not None):
            cgh = (cg1 or 0) + (cg2 or 0)
        if diff is None and dz is not None and cgh is not None:
            diff = dz - cgh
        if dz is not None:
            n_dz += 1
        sup_vals = []
        for (col, sup, oid) in pairs:
            odd = num(v[col - 1]) if col - 1 < len(v) else None
            rcv = num(v[col]) if col < len(v) else None
            if odd is None and rcv is None:
                continue
            sup_vals.append((sup, oid, odd, rcv))
        if not sup_vals and dz is None and cgh is None:
            continue
        if sup_vals:
            for (sup, oid, odd, rcv) in sup_vals:
                purchase.append([sn, code, name, unit, dz, cg1, cg2, cgh, diff, sup, oid, odd, rcv])
                n += 1
        else:
            purchase.append([sn, code, name, unit, dz, cg1, cg2, cgh, diff, "", "", None, None])
            n += 1
    proj_meta.append([sn, len(pairs), n, n_dz])
    log(f"  [{sn}] 供应商块 {len(pairs)} 个 / 明细 {n} 行 / 其中有设计量 {n_dz} 行")
wb.close()
n_pur = w("项目采购到货明细", ["项目", "物料编码", "物料名称", "单位", "设计量", "采购1", "采购2", "采购汇总",
                     "设计与提采差额", "供应商", "订单号", "订单数量", "接收数量"], purchase)

# ==================== 4. 物料主数据候选 ====================
log("\n## 物料主数据候选")
key = {}
for tag, name, model, unit, *_ in stock:
    if not model:
        continue
    k = (model, unit)
    d = key.setdefault(k, {"型号": model, "单位": unit, "名称": "", "集成商": set(), "编码": set()})
    if name and not d["名称"]:
        d["名称"] = name
    d["集成商"].add(tag)
for p in purchase:
    code, name, unit = p[1], p[2], p[3]
    if not name:
        continue
    k = (name, unit)
    d = key.setdefault(k, {"型号": name, "单位": unit, "名称": name, "集成商": set(), "编码": set()})
    if code:
        d["编码"].add(code)
for tag, proj, site, model, qty, rmk in design:
    k = (model, "")
    d = key.setdefault(k, {"型号": model, "单位": "", "名称": "", "集成商": set(), "编码": set()})
    d["集成商"].add(tag)

mat = []
for i, (k, d) in enumerate(sorted(key.items()), start=1):
    mat.append([f"M{i:04d}", d["名称"] or d["型号"], d["型号"], d["单位"],
                "|".join(sorted(d["集成商"])), "|".join(sorted(x for x in d["编码"] if x))])
n_mat = w("物料主数据", ["标准编码", "标准名称", "标准型号", "单位", "出现于集成商", "采购表物料编码"], mat)

# ==================== 5. 项目档案 ====================
projs = {}
for p in purchase:
    projs.setdefault(p[0], set())
for tag, proj, site, model, qty, rmk in design:
    if proj:
        projs.setdefault(proj, set()).add(tag)
for r in inbound:
    if r[7]:
        projs.setdefault(r[7], set()).add(r[0])
for r in outbound:
    if r[5]:
        projs.setdefault(r[5], set()).add(r[0])
pj = [[k, "|".join(sorted(v)), ""] for k, v in sorted(projs.items())]
n_pj = w("项目档案", ["项目名称", "涉及集成商", "备注"], pj)

# ==================== 6. 站点档案 ====================
sites = {}
for tag, proj, site, model, qty, rmk in design:
    if site:
        sites.setdefault((site, proj), set()).add(tag)
for r in outbound:
    to = r[6]
    if to and to not in ("-", "无"):
        sites.setdefault((to, r[5]), set()).add(r[0])
st = [[s, p, "|".join(sorted(v)), ""] for (s, p), v in sorted(sites.items())]
n_site = w("站点档案", ["站点名称", "项目", "集成商", "任务编码"], st)

log("\n" + "=" * 50)
log("汇总：")
log(f"  入库明细 {n_in} / 出库明细 {n_out} / 退料 {n_ret}")
log(f"  集成商汇总 {n_stk} / 盘点 {n_inv}")
log(f"  {ROLE_C} 站点设计明细 {n_des}")
log(f"  项目采购到货明细 {n_pur}")
log(f"  物料主数据候选 {n_mat}")
log(f"  项目档案 {n_pj} / 站点档案 {n_site}")

p = os.path.join(OUT, "extract_report.txt")
with io.open(p, "w", encoding="utf-8") as f:
    f.write("\n".join(report))
print("\n".join(report[-40:]))
