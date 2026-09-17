# -*- coding: utf-8 -*-
"""扫描 4 个本地源文件的真实结构，输出紧凑报告。"""
import os, sys, io, json
from openpyxl import load_workbook

# 台账根目录：优先取环境变量 LEDGER_DIR，默认用仓库内 data/ 目录
_REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = os.environ.get('LEDGER_DIR') or os.path.join(_REPO, 'data')
FILES = {
    "采购": BASE + r"\采购与设计量与到货量.xlsx",
    "集成商C": BASE + r"\集成商\集成商C\集成商C库存(9月8日).xlsx",
    "集成商A": BASE + r"\集成商\集成商A\上海集成商A材料库存(9月1日）.xlsx",
    "集成商B": BASE + r"\集成商\集成商B\集成商B-材料库存2026-8-26.xlsx",
}

out = []


def cell(ws, r, c):
    v = ws.cell(row=r, column=c).value
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return v


def dump(ws, r0, r1, c0, c1, label):
    out.append(f"--- {label} (行{r0}-{r1} 列{c0}-{c1}) ---")
    for r in range(r0, r1 + 1):
        vals = []
        for c in range(c0, c1 + 1):
            v = cell(ws, r, c)
            vals.append("" if v == "" else str(v).replace("\n", "\\n"))
        # 去掉尾部空列
        while vals and vals[-1] == "":
            vals.pop()
        if any(v != "" for v in vals):
            out.append(f"  r{r}: " + " | ".join(vals))


for tag, path in FILES.items():
    out.append("=" * 70)
    out.append(f"# {tag}  {path}")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        out.append(f"  !! 打开失败: {e}")
        continue
    for name in wb.sheetnames:
        ws = wb[name]
        try:
            mr, mc = ws.max_row, ws.max_column
        except Exception:
            mr, mc = -1, -1
        out.append(f"  [{tag}] 表「{name}」 rows={mr} cols={mc}")
    wb.close()

# ============ 采购表：逐项目看供应商布局 ============
out.append("=" * 70)
out.append("# 采购表 · 各项目的表头与供应商横向布局")
wb = load_workbook(FILES["采购"], read_only=True, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    out.append(f"\n## 项目「{name}」")
    dump(ws, 1, 4, 1, 8, "左侧固定列")
    # 找第 1 行（表头行）里的供应商名，第 3 行里的 订单数量/接收数量
    row1 = [(c, cell(ws, 1, c)) for c in range(9, 140)]
    sups = [(c, v) for c, v in row1 if v != "" and str(v).strip() != ""]
    out.append(f"  供应商块（列号:名称）: {sups}")
    row3 = [(c, cell(ws, 3, c)) for c in range(5, 140)]
    kinds = [(c, v) for c, v in row3 if v != ""]
    out.append(f"  第3行(订单/接收标记)共{len(kinds)}个: {kinds[:12]} ...")
wb.close()

# ============ 集成商文件：关键子表表头 ============
for tag in ["集成商C", "集成商A", "集成商B"]:
    out.append("=" * 70)
    out.append(f"# {tag} · 关键子表表头")
    wb = load_workbook(FILES[tag], read_only=True, data_only=True)
    for name in wb.sheetnames:
        ws = wb[name]
        try:
            mr, mc = ws.max_row, ws.max_column
        except Exception:
            continue
        if mr is None or mr < 2:
            continue
        out.append(f"\n## {tag}「{name}」 rows={mr} cols={mc}")
        dump(ws, 1, 3, 1, 12, "前3行")
    wb.close()

report = "\n".join(out)
p = r"C:\Users\qazws\WorkBuddy\2026-09-09-09-19-37\material-system\out\profile.txt"
with io.open(p, "w", encoding="utf-8") as f:
    f.write(report)
print("written", p, len(report), "chars")
